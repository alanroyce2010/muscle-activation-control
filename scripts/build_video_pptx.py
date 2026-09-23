"""Create an editable PowerPoint companion to the computational video.

Run with python3 (requires python-pptx and numpy); simulation data are read
from the existing video cache. Uses FFmpeg for embedded movie assets.
"""
from pathlib import Path
import json
import subprocess
import numpy as np
from pptx import Presentation
from pptx.chart.data import XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
VIDEO = ROOT / "results/submission_video"
ASSETS = VIDEO / "pptx_assets"
OUTPUT = ROOT / "docs/icra2027/muscle_control_video.pptx"
BG, INK, MUTED = "FDF0D5", "003049", "405967"
COLORS = ["780000", "C1121F", "669BBC", "003049"]
MODES = ["fixed_naive", "blended", "tracking", "current_state"]
LABELS = ["Fixed branch", "Blended", "Sign-test branch", "Relinearize at state"]
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def text(s, x, y, w, h, value, size=18, color=INK, bold=False, center=False):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(8)
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    return box


def slide(title, sub, notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(BG)
    text(s, .5, .25, 12, .2, "MUSCLE ACTIVATION CONTROL  /  COMPUTATIONAL STUDY", 10, MUTED, True)
    text(s, .5, .68, 12.3, .58, title, 28, bold=True)
    text(s, .5, 1.36, 12.3, .48, sub, 14, MUTED)
    text(s, .5, 7.13, 11.5, .2, "Simulation study • Hill-type elbow model • ICRA 2027", 10, MUTED)
    text(s, 12.1, 7.13, .7, .2, str(len(prs.slides)), 10, MUTED, center=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def chart(s, rect, series, xlabel, ylabel, bounds=None, legend=True):
    data = XyChartData()
    for label, xx, yy, color in series:
        seq = data.add_series(label)
        for x, y in zip(xx, yy):
            seq.add_data_point(float(x), float(y))
    x, y, w, h = rect
    c = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
        Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    c.font.name = "Arial"
    c.font.size = Pt(11)
    c.font.color.rgb = RGBColor.from_string(INK)
    c.has_legend = legend
    if legend:
        c.legend.position = XL_LEGEND_POSITION.BOTTOM
        c.legend.include_in_layout = False
        c.legend.font.size = Pt(10)
    for i, (_, _, _, color) in enumerate(series):
        line = c.series[i].format.line
        line.color.rgb = RGBColor.from_string(color)
        line.width = Pt(2)
        if i:
            line.dash_style = [MSO_LINE_DASH_STYLE.SOLID, MSO_LINE_DASH_STYLE.DASH,
                               MSO_LINE_DASH_STYLE.ROUND_DOT, MSO_LINE_DASH_STYLE.DASH_DOT][i % 4]
    for ax, title in [(c.category_axis, xlabel), (c.value_axis, ylabel)]:
        ax.has_title = True
        ax.axis_title.text_frame.text = title
        ax.axis_title.text_frame.paragraphs[0].font.size = Pt(12)
        ax.tick_labels.font.size = Pt(10)
    if bounds:
        c.category_axis.minimum_scale, c.category_axis.maximum_scale = bounds[:2]
        c.value_axis.minimum_scale, c.value_axis.maximum_scale = bounds[2:]
    return c


def movie(s, path, rect, autoplay=True):
    poster = ASSETS / (path.stem + "_poster.png")
    subprocess.run(["ffmpeg", "-v", "error", "-ss", "0.1", "-i", str(path),
                    "-frames:v", "1", "-y", str(poster)], check=True)
    x, y, w, h = rect
    shape = s.shapes.add_movie(str(path), Inches(x), Inches(y), Inches(w), Inches(h),
                              poster_frame_image=str(poster), mime_type="video/mp4")
    if autoplay:
        # python-pptx creates click-to-play timing; zero delay starts all arms together.
        for cond in s._element.xpath(".//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond"):
            cond.set("delay", "0")
    return shape


def clip(mode, name, start, duration):
    dest = ASSETS / f"{mode}_{name}.mp4"
    subprocess.run(["ffmpeg", "-v", "error", "-ss", str(start), "-i",
        str(VIDEO / "arm_clips" / f"{mode}.mp4"), "-t", str(duration),
        "-vf", "crop=720:650:0:0", "-c:v", "libx264", "-crf", "20", "-pix_fmt",
        "yuv420p", "-an", "-movflags", "+faststart", "-y", str(dest)], check=True)
    return dest


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    with np.load(VIDEO / "trajectories.npz") as z:
        d = dict(z)
    metrics = json.loads((VIDEO / "metrics.json").read_text())
    notes = ("Titles, captions, and charts are editable native PowerPoint objects. "
             "Right-click a chart and choose Edit Data to edit its workbook. "
             "Arm animations are embedded MP4 files; their frames are not editable shapes. "
             "Controller movies are configured to start together on entering a slide. "
             "Use Slide Show in desktop PowerPoint for media playback. "
             "The magnified arm angle is 60° + 12 × (actual angle − 60°); target uses the same transform. "
             "All chart values and metrics are actual physical degrees. "
             "Playback may differ in other presentation applications.")
    s = slide("One equilibrium. Multiple local models.",
              "How switching muscle dynamics change prediction and feedback control", notes)
    movie(s, VIDEO / "arm_clips/current_state.mp4", (.6, 1.95, 4.6, 4.6))
    text(s, 5.7, 2.3, 6.9, .8, "A simulated muscle-driven elbow", 25, bold=True)
    text(s, 5.7, 3.25, 6.8, 2.6,
         "Neural excitation → activation → joint motion\nNonlinear Hill-type muscle + gravity\nSame plant and task for every controller", 21)
    text(s, 5.7, 6.15, 6.8, .5, "Does the choice of local linear model matter?", 20, "780000", True)

    # Constitutive curves come from the existing scientific model, not fitted drawings.
    code = """import sys,json,numpy as np
sys.path.insert(0,'src')
from muscle_activation_control import muscle,linearize
a=linearize.find_equilibrium(np.deg2rad(60))
x=np.linspace(-.035,.035,101);v=np.linspace(-.5,.5,101)
print(json.dumps({'x':x.tolist(),'adot':muscle.activation_derivative(a,a+x).tolist(),
'v':v.tolist(),'force':(muscle.contractile_force(a,1.,v)/a).tolist()}))
"""
    curves = json.loads(subprocess.check_output([str(ROOT / ".venv/bin/python"), "-c", code], cwd=ROOT))
    s = slide("At rest, both switches meet.", "Two switching surfaces → four candidate branch linearizations.", notes)
    chart(s, (.55, 2.05, 5.9, 3.75), [("Activation", curves["x"], curves["adot"], COLORS[1])],
          "Excitation − activation, u − a", "Activation rate (1/s)", legend=False)
    chart(s, (6.85, 2.05, 5.9, 3.75), [("Force–velocity", curves["v"], curves["force"], INK)],
          "Normalized fibre velocity (lengths/s)", "Active force / activation", legend=False)
    text(s, .7, 6.1, 5.6, .5, "u = a at equilibrium", 23, center=True)
    text(s, 6.9, 6.1, 5.6, .5, "v = 0 at equilibrium", 23, center=True)

    s = slide("A small input exposes the mismatch.",
              "Open-loop prediction • excitation step +0.01 • initial equilibrium 60°", notes)
    tt = np.arange(51)*10
    series = [(label, tt, np.rad2deg(d[key][:, 0])-60, color) for key, label, color in
              [("pred_true", "Nonlinear simulation", "780000"),
               ("pred_matched", "Matched branches", INK), ("pred_wrong", "Both wrong", "C1121F")]]
    chart(s, (.6, 2., 8.5, 4.7), series, "Simulation time (ms)", "Angle change (deg)")
    text(s, 9.5, 3.0, 3.2, 1.0, f"{metrics['prediction']['wrong_error_deg']:.3f}°", 38, "C1121F", True)
    text(s, 9.5, 4.05, 3.2, 1.7, "Wrong-branch error\nat 500 ms\nSame state and excitation", 19)

    segments = [("raise", .10, .75, "Raising: the blended model overshoots.", "step_up"),
                ("lower", 1.15, 1.8, "Lowering: the linearization point matters.", "step_down"),
                ("track", 2.3, 4.3, "Tracking: compare the full motion.", "sinusoid")]
    for name, t0, t1, title, metric in segments:
        print(f"Building {name} slide", flush=True)
        s = slide(title, "Identical plant and MPC weights • 10 ms control interval • 200 ms horizon", notes +
                  " Reported metrics use the complete original evaluation windows, not just the displayed clip.")
        text(s, .5, 1.87, 12.3, .3,
             "ARM MOTION ×12 ABOUT 60°  •  DASHED: TARGET  •  PLOTS AND METRICS: ACTUAL DEGREES",
             12, "780000", True, True)
        for i, (mode, label, color) in enumerate(zip(MODES, LABELS, COLORS)):
            left = .52 + 3.16*i
            text(s, left, 2.23, 2.9, .35, label, 17, INK if i == 2 else color, True, True)
            path = clip(mode, name, t0*10, (t1-t0)*10)
            movie(s, path, (left+.11, 2.64, 2.68, 2.42))
            key = "rms_track_deg" if name == "track" else "overshoot_deg"
            val = metrics["controllers"][mode][metric][key]
            desc = "RMS error" if name == "track" else "Overshoot"
            text(s, left, 5.08, 2.9, .35, f"{desc}: {val:.3f}°", 16, INK, center=True)
        t = d["fixed_naive_t"]
        mask = (t >= t0) & (t <= t1)
        series = [("Target", t[mask], np.rad2deg(d['fixed_naive_theta_ref'][mask]), MUTED)]
        series += [(label, t[mask], np.rad2deg(d[f"{mode}_x"][mask, 0]), col)
                   for mode, label, col in zip(MODES, LABELS, COLORS)]
        chart(s, (.7, 5.53, 11.9, 1.48), series, "Simulation time (s)", "Angle (deg)", legend=False)

    s = slide("Co-contraction reduces the velocity kink.",
              "Brachialis + triceps • equilibrium sweep at a fixed 60° posture", notes)
    movie(s, VIDEO / "arm_clips/antagonist_cocontraction.mp4", (.55, 2., 4.4, 4.4))
    chart(s, (5.25, 2.05, 7.45, 4.05), [("Slope ratio", d["pair_at"], d["pair_ratio"], INK),
          ("Equal slopes", [0,.5], [1,1], "C1121F")], "Triceps equilibrium activation", "Velocity slope ratio")
    text(s, 5.45, 6.25, 7.15, .6, "Cancellation depends on geometry; activation switches remain.", 18)

    s = slide("Read the branch. Update the operating point.",
              "A computational result for linearization-based muscle control", notes)
    for y, num, title, body in [(2.15, "01", "Equilibrium sits on switching surfaces.", "A blended Jacobian hides the branch structure."),
        (3.65, "02", "Prediction mismatch survives as transient error.", "Same plant, weights and horizon across all four controllers."),
        (5.15, "03", "Relinearize at the measured state.", "Use explicit branches; model transport delay.")]:
        text(s, .65, y, .75, .6, num, 30, "C1121F", True)
        text(s, 1.65, y, 10.9, .6, title, 25, bold=True)
        text(s, 1.65, y+.66, 10.9, .55, body, 18, MUTED)
    text(s, 1.65, 6.65, 10.9, .3, "Baseline weighting shown • hardware validation remains future work", 13, MUTED)

    s = slide("Complete video", "Original 2:30 animation • click the video to play", notes)
    movie(s, VIDEO / "icra2027_muscle_control_light.mp4", (2.05, 1.96, 9.25, 5.203), autoplay=False)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Saved {OUTPUT} ({OUTPUT.stat().st_size/1e6:.1f} MB)", flush=True)


if __name__ == "__main__":
    main()
