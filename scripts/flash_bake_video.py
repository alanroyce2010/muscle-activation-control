"""Bake the editable 4-slide flash deck into a one-slide, one-video deck.

Why: PowerPoint's timed slide advance (advTm) is unreliable on slides whose
clip starts automatically (measured 2026-09-21: with the media command's
duration set to 1 ms the slide never advances; with it set to the clip length
it advances early and non-deterministically). A single 64 s video that starts
on slide entry needs no clicks and no slide-advance logic.

Steps: export the editable deck to PDF through desktop PowerPoint (AppleScript),
rasterize each slide to 1920x1080 as the background, overlay each slide's clip
at the exact rectangle of its movie shape, hold the last frame up to the pitch
cue, concatenate, and embed the result full-bleed on one slide with the
start-automatically timing from flash_slides_video.movie().

Inputs:  docs/poster/05_AlanRoyceGabriel_Samuel_4slides_editable.pptx
         (edit slide text there; the clips inside it are the ones overlaid)
Outputs: results/submission_video/flash_clips/flash_60s.mp4
         docs/poster/05_AlanRoyceGabriel_Samuel.pptx  (upload this one)
         docs/poster/05_AlanRoyceGabriel_Samuel.pdf   (check copy)

Segment lengths (seconds) follow docs/poster/flash_pitch.md: 35, 9, 13, 7.
Requires macOS with Microsoft PowerPoint, ffmpeg, pdftoppm, python-pptx.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import flash_slides_video as F  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

ROOT = F.ROOT
EDITABLE = ROOT / "docs/poster/05_AlanRoyceGabriel_Samuel_4slides_editable.pptx"
OUT = ROOT / "docs/poster/05_AlanRoyceGabriel_Samuel.pptx"
OUT_PDF = OUT.with_suffix(".pdf")
VIDEO = F.CLIPS / "flash_60s.mp4"
SEGMENT_S = [35.0, 9.0, 13.0, 7.0]
W, H = 1920, 1080
EMU_PER_PX = 13.333 * 914400 / W


def run(cmd, **kw):
    subprocess.run(cmd, check=True, **kw)


def powerpoint_pdf(pptx: Path, pdf: Path):
    """Open a copy in PowerPoint (never the user's own path) and save as PDF."""
    script = f'''
    set inFile to POSIX file "{pptx}"
    tell application "Microsoft PowerPoint"
      open inFile
      delay 5
      set pres to presentation "{pptx.name}"
      save pres in (POSIX file "{pdf}") as save as PDF
      delay 2
      close pres saving no
    end tell'''
    run(["osascript", "-e", script])


def movie_rects(pptx: Path):
    rects = []
    with zipfile.ZipFile(pptx) as z:
        i = 1
        while f"ppt/slides/slide{i}.xml" in z.namelist():
            x = z.read(f"ppt/slides/slide{i}.xml").decode()
            m = re.search(r'<p:pic>.*?<a:videoFile.*?<a:off x="(\d+)" y="(\d+)"/>'
                          r'<a:ext cx="(\d+)" cy="(\d+)"/>', x, flags=re.S)
            rects.append([int(round(int(v) / EMU_PER_PX)) for v in m.groups()])
            rel = z.read(f"ppt/slides/_rels/slide{i}.xml.rels").decode()
            media = re.search(r'Target="\.\./media/(media\d+\.mp4)"', rel).group(1)
            rects[-1].append(media)
            i += 1
    return rects


def main():
    tmp = Path(tempfile.mkdtemp(prefix="flash_bake_"))
    src = tmp / "editable_copy.pptx"
    shutil.copy(EDITABLE, src)
    powerpoint_pdf(src, tmp / "deck.pdf")
    run(["pdftoppm", "-png", "-r", "144", "-scale-to-x", str(W), "-scale-to-y", str(H),
         str(tmp / "deck.pdf"), str(tmp / "bg")])
    bgs = sorted(tmp.glob("bg-*.png"))
    rects = movie_rects(src)
    assert len(bgs) == len(rects) == len(SEGMENT_S), (len(bgs), len(rects))
    with zipfile.ZipFile(src) as z:
        for _, _, _, _, media in rects:
            (tmp / media).write_bytes(z.read(f"ppt/media/{media}"))
    segs = []
    for k, (bg, (x, y, w, h, media), total) in enumerate(zip(bgs, rects, SEGMENT_S), 1):
        clip = tmp / media
        hold = max(0.0, total - F.duration(clip))
        seg = tmp / f"seg{k}.mp4"
        # even dimensions for yuv420p
        w, h = w - w % 2, h - h % 2
        run(["ffmpeg", "-v", "error", "-y", "-loop", "1", "-framerate", "30", "-i", str(bg),
             "-i", str(clip), "-filter_complex",
             f"[1:v]scale={w}:{h}:flags=lanczos,tpad=stop_mode=clone:stop_duration={hold:.3f}[c];"
             f"[0:v][c]overlay={x}:{y}:shortest=1,trim=duration={total},format=yuv420p",
             "-r", "30", "-c:v", "libx264", "-crf", "20", "-preset", "slow", "-an", str(seg)])
        segs.append(seg)
        print(f"segment {k}: {media} at ({x},{y}) {w}x{h}, hold {hold:.1f} s, total {total} s")
    lst = tmp / "concat.txt"
    lst.write_text("".join(f"file '{s}'\n" for s in segs))
    run(["ffmpeg", "-v", "error", "-y", "-f", "concat", "-safe", "0", "-i", str(lst),
         "-c:v", "libx264", "-crf", "20", "-preset", "slow", "-pix_fmt", "yuv420p",
         "-movflags", "+faststart", "-an", str(VIDEO)])
    print(f"video {VIDEO.name}: {F.duration(VIDEO):.1f} s")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = F.BG
    F.movie(s, VIDEO, 0, 0, 13.333, 7.5)
    cues = []
    t = 0.0
    for k, d in enumerate(SEGMENT_S, 1):
        cues.append(f"{t:.0f} s segment {k}")
        t += d
    s.notes_slide.notes_text_frame.text = (
        f"One slide, one {t:.0f} s video, starts automatically; nothing to click. "
        "Cues (video time): " + "; ".join(cues) + ". Slide-1 clip runs at 1.5x and holds "
        "its last frame until the 35 s cue. Script: docs/poster/flash_pitch.md. If the "
        "video does not start by itself (non-PowerPoint player), one click on it starts "
        "it. Edit text in the 4-slide editable deck and rerun scripts/flash_bake_video.py.")
    prs.save(OUT)
    print("saved", OUT, f"{OUT.stat().st_size / 1e6:.1f} MB")
    shutil.copy(OUT, tmp / "check.pptx")
    powerpoint_pdf(tmp / "check.pptx", OUT_PDF)
    print("saved", OUT_PDF)
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
